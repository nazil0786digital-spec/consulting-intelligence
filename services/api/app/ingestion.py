import hashlib
import re
from io import BytesIO
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.embeddings import EmbeddingProvider, get_embedding_provider
from app.models import Document, DocumentChunk


MAX_DOCUMENT_BYTES = 1_000_000
CHUNK_WORDS = 180
SUPPORTED_CONTENT_TYPES = {"text/plain", "application/pdf", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/octet-stream"}


def content_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def chunk_text(text: str, words_per_chunk: int = CHUNK_WORDS) -> list[str]:
    """Create stable, readable chunks for the first ingestion release."""
    words = re.findall(r"\S+", text)
    return [" ".join(words[index : index + words_per_chunk]) for index in range(0, len(words), words_per_chunk)]


def extract_document_sections(*, title: str, content_type: str | None, content: bytes) -> tuple[str, list[tuple[str, str]]]:
    """Extract text locally and preserve a citation-friendly source locator per section."""
    suffix = Path(title).suffix.lower()
    if suffix == ".txt" or content_type == "text/plain":
        try:
            return "text", [(content.decode("utf-8"), "text")]
        except UnicodeDecodeError as error:
            raise ValueError("Plain-text documents must use UTF-8 encoding.") from error
    if suffix == ".pdf" or content_type == "application/pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(content))
            sections = [(page.extract_text() or "", f"page:{index + 1}") for index, page in enumerate(reader.pages)]
        except Exception as error:
            raise ValueError("The PDF could not be read. Password-protected or malformed PDFs are not supported.") from error
        return "pdf", sections
    if suffix == ".docx" or content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        try:
            from docx import Document as WordDocument
            word_document = WordDocument(BytesIO(content))
            sections = [(paragraph.text, f"paragraph:{index + 1}") for index, paragraph in enumerate(word_document.paragraphs)]
        except Exception as error:
            raise ValueError("The Word document could not be read. Upload a valid .docx file.") from error
        return "docx", sections
    raise ValueError("Supported documents are UTF-8 .txt, text-based .pdf, and .docx files.")


def ingest_document(
    session: Session,
    *,
    organization_id: str,
    title: str,
    source_type: str,
    content_type: str | None,
    content: bytes,
    embedding_provider: EmbeddingProvider | None = None,
) -> Document:
    if not content:
        raise ValueError("The uploaded document is empty.")
    if len(content) > MAX_DOCUMENT_BYTES:
        raise ValueError("The uploaded document exceeds the 1 MB Phase 2 limit.")
    document_format, sections = extract_document_sections(title=title, content_type=content_type, content=content)
    digest = content_hash(content)
    duplicate = session.scalar(select(Document).where(Document.organization_id == organization_id, Document.integrity_hash == digest))
    if duplicate:
        raise ValueError("This exact document has already been indexed for the organization.")

    chunk_rows = [
        (chunk, locator)
        for section, locator in sections
        for chunk in chunk_text(section)
    ]
    if not chunk_rows:
        raise ValueError("The uploaded document contains no indexable text. Scanned PDFs need OCR, which is not enabled yet.")

    provider = embedding_provider or get_embedding_provider()
    document = Document(
        organization_id=organization_id,
        title=title,
        source_type=source_type,
        status="processing",
        integrity_hash=digest,
        metadata_json={"format": document_format, "content_type": content_type, "bytes": len(content)},
    )
    session.add(document)
    session.flush()
    session.add_all(
        [
            DocumentChunk(
                organization_id=organization_id,
                document_id=document.id,
                content=chunk,
                source_locator=locator,
                embedding=provider.embed(chunk),
                metadata_json={"source_locator": locator, "word_count": len(chunk.split())},
            )
            for chunk, locator in chunk_rows
        ]
    )
    document.status = "indexed"
    session.commit()
    session.refresh(document)
    return document


def ingest_text_document(
    session: Session,
    *,
    organization_id: str,
    title: str,
    source_type: str,
    content: bytes,
    embedding_provider: EmbeddingProvider | None = None,
) -> Document:
    return ingest_document(
        session,
        organization_id=organization_id,
        title=title,
        source_type=source_type,
        content_type="text/plain",
        content=content,
        embedding_provider=embedding_provider,
    )
